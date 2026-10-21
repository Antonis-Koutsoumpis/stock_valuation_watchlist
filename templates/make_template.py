"""Builds templates/manual_entries_template.xlsx, the fixed layout for manual data.
Upload it to Google Drive and open it as a Google Sheet (or keep it as Excel)."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from rule1.manual import TABS

OUT = Path(__file__).resolve().parent / "manual_entries_template.xlsx"

METRICS = [
    ("revenue", "Revenue", "statements"),
    ("net_income", "Net income", "statements"),
    ("eps_diluted", "EPS (diluted)", "statements"),
    ("operating_income", "Operating income (EBIT), used for NOPAT", "statements"),
    ("pretax_income", "Pre-tax income, used for the tax rate", "statements"),
    ("tax_provision", "Income tax, used for the tax rate", "statements"),
    ("total_debt", "Debt (short + long term)", "statements"),
    ("cash", "Cash and cash equivalents", "statements"),
    ("equity", "Shareholders' equity", "statements"),
    ("shares_outstanding", "Shares outstanding", "statements"),
    ("operating_cash_flow", "Operating cash flow", "statements"),
    ("capex", "Capital expenditure (positive number)", "statements"),
    ("dividends_paid", "Dividends paid (positive number)", "statements"),
    ("roic", "ROIC, yearly (overrides the calculation)", "statements"),
    ("roic_ttm", "ROIC, TTM per quarter (overrides the calculation)", "statements"),
    ("bvps", "BVPS (overrides the calculation)", "statements"),
    ("fcf", "Free cash flow (overrides OCF - CapEx)", "statements"),
    ("close", "Closing price", "daily"),
    ("volume", "Volume", "daily"),
    ("trailing_pe", "Trailing P/E", "daily"),
    ("forward_pe", "Forward P/E", "daily"),
    ("peg_5y", "PEG ratio (5yr expected)", "daily"),
    ("price_sales", "Price/Sales", "daily"),
    ("price_book", "Price/Book", "daily"),
    ("trailing_div_yield", "Trailing annual dividend yield (0.012 = 1.2%)", "daily"),
    ("forward_div_yield", "Forward annual dividend yield (0.012 = 1.2%)", "daily"),
    ("payout_ratio", "Payout ratio (0.3 = 30%)", "daily"),
]

HELP = [
    "How to fill in this sheet",
    "",
    "statements: one row per value. period = 2026Q1 for a quarter (calendar quarter) or 2025 for a year.",
    "  Values in the stock's own currency, full numbers (12500000000, not 12.5B).",
    "  Quarter values are for that quarter only, not year-to-date.",
    "daily: rarely needed. date = YYYY-MM-DD.",
    "dcf: one row per stock per quarter with the 7 site estimates (leave a cell empty if a site has none).",
    "  middle = median or mean: which one the tool uses as the middle of the DCF range.",
    "metric: pick from the dropdown; the list with descriptions is on the 'metrics' tab.",
    "ticker: as on the watchlist (AMZN, FTK).",
    "Manual values always win over automatically extracted ones.",
    "Don't rename tabs or column headers; the tool reads them by name.",
]

HEADER = Font(bold=True, color="FFFFFF")
FILL = PatternFill("solid", fgColor="305496")


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "how_to"
    for i, line in enumerate(HELP, 1):
        ws.cell(i, 1, line).font = Font(bold=i == 1)
    ws.column_dimensions["A"].width = 110

    ref = wb.create_sheet("metrics")
    ref.append(["metric", "description", "tab"])
    for row in METRICS:
        ref.append(list(row))

    for tab, cols in TABS.items():
        sheet = wb.create_sheet(tab)
        sheet.append(cols)
        for i, _ in enumerate(cols, 1):
            sheet.cell(1, i).font, sheet.cell(1, i).fill = HEADER, FILL
            sheet.column_dimensions[sheet.cell(1, i).column_letter].width = 16
        sheet.freeze_panes = "A2"
        if "metric" in cols:
            col = sheet.cell(1, cols.index("metric") + 1).column_letter
            dv = DataValidation(type="list", formula1=f"=metrics!$A$2:$A${len(METRICS) + 1}", allow_blank=True)
            sheet.add_data_validation(dv)
            dv.add(f"{col}2:{col}5000")
        if tab == "dcf":
            dv = DataValidation(type="list", formula1='"median,mean"', allow_blank=True)
            sheet.add_data_validation(dv)
            dv.add("J2:J5000")

    ex = wb.create_sheet("examples")
    ex.append(["Made-up examples showing the layout; the tool ignores this tab."])
    ex.append(TABS["statements"])
    ex.append(["FTK", "2020Q4", "revenue", 123000000, "from DeGiro"])
    ex.append(["FTK", "2019", "total_debt", 45000000, "a yearly value"])
    ex.append([])
    ex.append(TABS["dcf"])
    ex.append(["AMZN", "2026Q4", 100, 110, 120, 130, 140, 150, 160, "median", ""])
    wb.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
